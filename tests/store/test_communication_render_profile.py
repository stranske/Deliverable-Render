"""End-to-end synthesis profile: one validated store, three public commands."""

from __future__ import annotations

import json
import shutil
import subprocess
from copy import deepcopy
from pathlib import Path

import pytest
from pptx import Presentation

from deliverable_render.cli.render_docx import main as docx_main
from deliverable_render.cli.render_html import main as html_main
from deliverable_render.cli.render_pptx import main as pptx_main
from deliverable_render.store.communication import (
    CommunicationRenderError,
    adapt_profile,
    adapt_store,
    load_profile,
)
from deliverable_render.store.validate import validate_store
from docx import Document

HERE = Path(__file__).parent.parent
FIXTURES = HERE / "fixtures" / "stores"
STORE = FIXTURES / "communication_render.json"
PATHS = FIXTURES / "communication_paths.json"
MEMO = FIXTURES / "communication_memo.json"
MANIFEST = FIXTURES / "communication_deck.json"
TEMPLATE = HERE / "fixtures" / "deck" / "synthetic_template.pptx"


def test_validated_store_renders_all_three_public_outputs(tmp_path: Path) -> None:
    assert validate_store(STORE).valid
    data, paths = load_profile(STORE, PATHS)
    adapted = adapt_store(data, paths)
    assert [record.record_id for record in adapted.records] == [
        "entry-0-mention-0",
        "entry-0-thesis",
        "entry-0-publication",
    ]
    assert adapted.records[0].evidence[0].page == 3
    assert adapted.records[1].text == "Assets of $36B"
    assert adapted.records[2].text == "Quarterly letter published"
    assert adapted.records[2].evidence[0].page == 4

    html = tmp_path / "hub.html"
    pptx = tmp_path / "deck.pptx"
    docx = tmp_path / "memo.docx"
    assert (
        html_main(["--store", str(STORE), "--document-paths", str(PATHS), "--out", str(html)]) == 0
    )
    assert (
        pptx_main(
            [
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--manifest",
                str(MANIFEST),
                "--template",
                str(TEMPLATE),
                "--out",
                str(pptx),
            ]
        )
        == 0
    )
    assert (
        docx_main(
            [
                "--profile",
                "communication-synthesis",
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--memo-overlay",
                str(MEMO),
                "--out",
                str(docx),
            ]
        )
        == 0
    )
    assert all(path.stat().st_size > 100 for path in (html, pptx, docx))
    page = html.read_text(encoding="utf-8")
    assert "file:///synthetic/source%20document.pdf#page=3" in page
    assert "https://" not in page and "http://" not in page
    assert "Capital of $36B" in page
    assert "Quarterly letter published" in page
    assert "file:///synthetic/source%20document.pdf#page=4" in page
    assert Presentation(pptx).slides[0].shapes.title.text == "Capital"
    assert "Assets of $36B" in " ".join(p.text for p in Document(docx).paragraphs)


def test_installed_console_commands_render_all_three_outputs(tmp_path: Path) -> None:
    outputs = {
        "render-html-hub": tmp_path / "console-hub.html",
        "render-pptx-deck": tmp_path / "console-deck.pptx",
        "render-docx-memo": tmp_path / "console-memo.docx",
    }
    arguments = {
        "render-html-hub": [
            "--store",
            str(STORE),
            "--document-paths",
            str(PATHS),
        ],
        "render-pptx-deck": [
            "--store",
            str(STORE),
            "--document-paths",
            str(PATHS),
            "--manifest",
            str(MANIFEST),
            "--template",
            str(TEMPLATE),
        ],
        "render-docx-memo": [
            "--profile",
            "communication-synthesis",
            "--store",
            str(STORE),
            "--document-paths",
            str(PATHS),
            "--memo-overlay",
            str(MEMO),
        ],
    }

    for command, output in outputs.items():
        executable = shutil.which(command)
        assert executable is not None, f"console script is not installed: {command}"
        result = subprocess.run(
            [executable, *arguments[command], "--out", str(output)],
            capture_output=True,
            text=True,
            check=False,
        )
        assert result.returncode == 0, result.stdout + result.stderr
        assert output.stat().st_size > 100

    page = outputs["render-html-hub"].read_text(encoding="utf-8")
    assert "https://" not in page and "http://" not in page


def test_missing_document_path_fails_before_output(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    mapping = tmp_path / "paths.json"
    mapping.write_text("{}", encoding="utf-8")
    commands = (
        (
            html_main,
            ["--store", str(STORE), "--document-paths", str(mapping)],
            tmp_path / "hub.html",
        ),
        (
            pptx_main,
            [
                "--store",
                str(STORE),
                "--document-paths",
                str(mapping),
                "--manifest",
                str(MANIFEST),
                "--template",
                str(TEMPLATE),
            ],
            tmp_path / "deck.pptx",
        ),
        (
            docx_main,
            [
                "--profile",
                "communication-synthesis",
                "--store",
                str(STORE),
                "--document-paths",
                str(mapping),
                "--memo-overlay",
                str(MEMO),
            ],
            tmp_path / "memo.docx",
        ),
    )
    for command, args, out in commands:
        assert command([*args, "--out", str(out)]) == 2
        assert not out.exists()

    errors = capsys.readouterr().err
    assert errors.count("document-path mapping missing") == len(commands)
    data = json.loads(STORE.read_text(encoding="utf-8"))
    with pytest.raises(CommunicationRenderError, match="must be an absolute local path"):
        adapt_store(data, {"doc-1": "relative/source.pdf"})


def test_adapter_requires_mapping_for_uncited_documents(tmp_path: Path) -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    data["documents"].append({"name": "uncited-doc", "date": "2025-02-01"})
    store_path = tmp_path / "store.json"
    store_path.write_text(json.dumps(data), encoding="utf-8")
    assert validate_store(store_path).valid

    with pytest.raises(
        CommunicationRenderError, match="document-path mapping missing for 'uncited-doc'"
    ):
        adapt_profile(store_path, PATHS)


@pytest.mark.parametrize("pointer_path", [("mentions", 0, "src"), ("pub", "src")])
def test_validator_rejects_evidence_source_not_in_documents(
    tmp_path: Path, pointer_path: tuple[str | int, ...]
) -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    pointer = data["entries"][0]
    for part in pointer_path:
        pointer = pointer[part]
    pointer["stable_id"] = "ghost-doc"
    store_path = tmp_path / f"orphan-evidence-{pointer_path[0]}.json"
    store_path.write_text(json.dumps(data), encoding="utf-8")

    report = validate_store(store_path)

    assert not report.valid
    assert any(
        issue.code == "orphan-reference"
        and issue.path.endswith("/src/source_id")
        and "'ghost-doc' does not resolve in documents" in issue.message
        for issue in report.issues
    )


def test_validator_rejects_duplicate_document_identity(tmp_path: Path) -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    data["documents"].append({"name": "doc-1"})
    store_path = tmp_path / "duplicate-document.json"
    store_path.write_text(json.dumps(data), encoding="utf-8")

    report = validate_store(store_path)

    assert not report.valid
    assert any(
        issue.code == "duplicate-document-identity"
        and issue.path == "/documents/1/name"
        and "duplicate document identity 'doc-1'" in issue.message
        for issue in report.issues
    )


def test_validator_rejects_duplicate_document_stable_id(tmp_path: Path) -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    data["documents"][0]["stable_id"] = "doc-1"
    data["documents"].append({"name": "doc-2", "stable_id": "doc-1"})
    store_path = tmp_path / "duplicate-document-stable-id.json"
    store_path.write_text(json.dumps(data), encoding="utf-8")

    report = validate_store(store_path)

    assert not report.valid
    assert any(
        issue.code == "duplicate-document-identity"
        and issue.path == "/documents/1/stable_id"
        and "duplicate document identity 'doc-1'" in issue.message
        for issue in report.issues
    )


def test_validator_prefers_distinct_stable_ids_over_equal_names(tmp_path: Path) -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    data["documents"][0]["stable_id"] = "doc-1"
    data["documents"].append({"name": "doc-1", "stable_id": "doc-2"})
    store_path = tmp_path / "same-name-distinct-stable-ids.json"
    store_path.write_text(json.dumps(data), encoding="utf-8")

    report = validate_store(store_path)

    assert report.valid


def test_adapter_does_not_replace_an_explicit_blank_stable_id_with_name() -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    data["documents"][0]["stable_id"] = ""

    with pytest.raises(CommunicationRenderError, match=r"documents\[0\]\.stable_id"):
        adapt_store(data, json.loads(PATHS.read_text(encoding="utf-8")))


@pytest.mark.parametrize("pointer_path", [("mentions", 0, "src"), ("pub", "src")])
@pytest.mark.parametrize("page_location", ["page", "locator"])
@pytest.mark.parametrize("page", [0, 3.0, None])
def test_validator_rejects_invalid_evidence_page_values(
    tmp_path: Path,
    pointer_path: tuple[str | int, ...],
    page_location: str,
    page: int | float | None,
) -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    pointer = data["entries"][0]
    for part in pointer_path:
        pointer = pointer[part]
    pointer.pop("page", None)
    if page_location == "locator":
        pointer["locator"] = {} if page is None else {"page": page}
    elif page is not None:
        pointer["page"] = page
    store_path = tmp_path / "invalid-page.json"
    store_path.write_text(json.dumps(data), encoding="utf-8")

    report = validate_store(store_path)
    pointer_json_path = (
        "/entries/0/mentions/0/src" if pointer_path[0] == "mentions" else "/entries/0/pub/src"
    )
    page_json_path = pointer_json_path + (
        "/locator/page" if page_location == "locator" else "/page"
    )
    assert not report.valid
    assert any(
        issue.code == "invalid-evidence-page"
        and issue.path == page_json_path
        and issue.message == "positive one-based evidence page required"
        for issue in report.issues
    )


def test_missing_page_fails_validation_before_projection(tmp_path: Path) -> None:
    for pointer_path in (("mentions", 0, "src"), ("pub", "src")):
        data = json.loads(STORE.read_text(encoding="utf-8"))
        pointer = data["entries"][0]
        for part in pointer_path:
            pointer = pointer[part]
        del pointer["page"]
        without_page = tmp_path / f"missing-page-{pointer_path[0]}.json"
        without_page.write_text(json.dumps(data), encoding="utf-8")
        assert not validate_store(without_page).valid
        out = tmp_path / f"hub-{pointer_path[0]}.html"
        assert (
            html_main(
                ["--store", str(without_page), "--document-paths", str(PATHS), "--out", str(out)]
            )
            == 2
        )
        assert not out.exists()


def test_blank_unsourced_mention_is_skipped(tmp_path: Path) -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    data["entries"] = [
        {
            "id": "solo",
            "name": "Solo",
            "first": "2025-Q1",
            "mentions": [{"q": "2025-Q1", "text": "   "}],
        }
    ]
    with pytest.raises(CommunicationRenderError, match="no renderable entry text"):
        adapt_store(data, json.loads(PATHS.read_text(encoding="utf-8")))


def test_non_object_manifest_exits_with_diagnostic(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    manifest = tmp_path / "array-manifest.json"
    manifest.write_text("[]", encoding="utf-8")
    out = tmp_path / "deck.pptx"
    assert (
        pptx_main(
            [
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--manifest",
                str(manifest),
                "--template",
                str(TEMPLATE),
                "--out",
                str(out),
            ]
        )
        == 2
    )
    assert not out.exists()
    assert "manifest must be a JSON object" in capsys.readouterr().err


@pytest.mark.parametrize("template_contents", [None, b"not a PowerPoint package"])
def test_invalid_template_exits_with_diagnostic_without_output(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    template_contents: bytes | None,
) -> None:
    template = tmp_path / "invalid-template.pptx"
    if template_contents is not None:
        template.write_bytes(template_contents)
    out = tmp_path / "deck.pptx"

    assert (
        pptx_main(
            [
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--manifest",
                str(MANIFEST),
                "--template",
                str(template),
                "--out",
                str(out),
            ]
        )
        == 2
    )
    assert not out.exists()
    assert "render-pptx-deck: Package not found" in capsys.readouterr().err


def test_cross_entry_mention_is_not_silently_reassigned(tmp_path: Path) -> None:
    data = json.loads(STORE.read_text(encoding="utf-8"))
    second = deepcopy(data["entries"][0])
    second["id"] = "other"
    second["mentions"] = []
    data["entries"].append(second)
    data["entries"][0]["mentions"][0]["entry_id"] = "other"
    store_path = tmp_path / "cross-entry.json"
    store_path.write_text(json.dumps(data), encoding="utf-8")
    assert validate_store(store_path).valid
    with pytest.raises(CommunicationRenderError, match="names another entry"):
        adapt_store(data, json.loads(PATHS.read_text(encoding="utf-8")))


def test_memo_requires_reviewed_rows_and_distinct_periods(tmp_path: Path) -> None:
    out = tmp_path / "memo.docx"
    assert (
        docx_main(
            [
                "--profile",
                "communication-synthesis",
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--out",
                str(out),
            ]
        )
        == 2
    )
    assert not out.exists()
    overlay = json.loads(MEMO.read_text(encoding="utf-8"))
    overlay["period_prior"] = overlay["period_current"]
    path = tmp_path / "overlay.json"
    path.write_text(json.dumps(overlay), encoding="utf-8")
    assert (
        docx_main(
            [
                "--profile",
                "communication-synthesis",
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--memo-overlay",
                str(path),
                "--out",
                str(out),
            ]
        )
        == 2
    )
    assert not out.exists()


def test_docx_rejects_communication_options_with_legacy_profile(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "memo.docx"

    assert (
        docx_main(
            [
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--memo-overlay",
                str(MEMO),
                "--out",
                str(out),
            ]
        )
        == 2
    )
    assert not out.exists()
    error = capsys.readouterr().err
    assert "--profile communication-synthesis" in error
    assert "document-path mappings" in error


def test_no_material_change_and_existing_output_are_preserved(tmp_path: Path) -> None:
    overlay = json.loads(MEMO.read_text(encoding="utf-8"))
    overlay["changes"][0]["tier"] = "T3"
    path = tmp_path / "overlay.json"
    path.write_text(json.dumps(overlay), encoding="utf-8")
    docx = tmp_path / "memo.docx"
    assert (
        docx_main(
            [
                "--profile",
                "communication-synthesis",
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--memo-overlay",
                str(path),
                "--out",
                str(docx),
            ]
        )
        == 2
    )
    assert not docx.exists()

    html = tmp_path / "hub.html"
    html.write_text("existing", encoding="utf-8")
    args = ["--store", str(STORE), "--document-paths", str(PATHS), "--out", str(html)]
    assert html_main(args) == 2
    assert html.read_text(encoding="utf-8") == "existing"
    assert html_main(args + ["--force"]) == 0
    assert "Capital of $36B" in html.read_text(encoding="utf-8")


def test_force_cannot_replace_a_mapped_source_document(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(b"original source")
    mapping = tmp_path / "paths.json"
    mapping.write_text(json.dumps({"doc-1": str(source)}), encoding="utf-8")

    command_args = (
        (
            html_main,
            ["--store", str(STORE), "--document-paths", str(mapping)],
        ),
        (
            pptx_main,
            [
                "--store",
                str(STORE),
                "--document-paths",
                str(mapping),
                "--manifest",
                str(MANIFEST),
                "--template",
                str(TEMPLATE),
            ],
        ),
        (
            docx_main,
            [
                "--profile",
                "communication-synthesis",
                "--store",
                str(STORE),
                "--document-paths",
                str(mapping),
                "--memo-overlay",
                str(MEMO),
            ],
        ),
    )
    for command, args in command_args:
        assert command([*args, "--out", str(source), "--force"]) == 2
        assert source.read_bytes() == b"original source"

    errors = capsys.readouterr().err
    assert errors.count("output path aliases a rendering input") == 3


def test_force_cannot_replace_a_static_deck_asset(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    asset = tmp_path / "slide.txt"
    asset.write_text("Reviewed static slide copy", encoding="utf-8")
    template = tmp_path / "template.pptx"
    shutil.copyfile(TEMPLATE, template)
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "slides": [
                    {
                        "slide_id": "static-slide",
                        "title": "Static slide",
                        "layout": "Title and Content",
                        "source": "static:slide.txt",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )

    assert (
        pptx_main(
            [
                "--store",
                str(STORE),
                "--document-paths",
                str(PATHS),
                "--manifest",
                str(manifest),
                "--template",
                str(template),
                "--out",
                str(asset),
                "--force",
            ]
        )
        == 2
    )
    assert asset.read_text(encoding="utf-8") == "Reviewed static slide copy"
    assert "output path aliases a rendering input" in capsys.readouterr().err
