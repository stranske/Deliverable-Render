"""Manifest-gated deck builder completeness and store contract tests."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation

from deliverable_render.pptx.manifest import (
    DeckBuildError,
    DeckManifest,
    build_deck,
    enforce_completeness,
)
from deliverable_render.store import Store

FIXTURES = Path(__file__).parent / "fixtures" / "deck"
TEMPLATE = FIXTURES / "synthetic_template.pptx"
STORE = Store.from_json(Path(__file__).parent / "fixtures" / "synthetic_store.json")


def _load_manifest(name: str) -> DeckManifest:
    return DeckManifest.from_mapping(json.loads((FIXTURES / name).read_text(encoding="utf-8")))


def _slide_titles(pptx_bytes: bytes) -> list[str]:
    presentation = Presentation(BytesIO(pptx_bytes))
    titles: list[str] = []
    for slide in presentation.slides:
        if slide.shapes.title is not None and slide.shapes.title.text:
            titles.append(slide.shapes.title.text)
    return titles


def _slide_body_text(slide: object) -> str:
    shapes = getattr(slide, "shapes", None)
    if shapes is None:
        return ""
    title_shape = shapes.title
    for shape in shapes:
        if shape == title_shape:
            continue
        if hasattr(shape, "text") and shape.text:
            return str(shape.text)
    return ""


def _slide_bodies(pptx_bytes: bytes) -> list[str]:
    presentation = Presentation(BytesIO(pptx_bytes))
    return [_slide_body_text(slide) for slide in presentation.slides]


def test_missing_successor_without_dropped_reason_fails_and_names_slide_id() -> None:
    prior = _load_manifest("manifest_v1.json")
    current = _load_manifest("manifest_v2_incomplete.json")

    with pytest.raises(DeckBuildError, match="slide_id 'risk'"):
        build_deck(TEMPLATE, current, STORE, prior=prior)


def test_dropped_reason_allows_build_and_counts_summary() -> None:
    prior = _load_manifest("manifest_v1.json")
    current = _load_manifest("manifest_v2_complete.json")

    pptx_bytes, summary = build_deck(TEMPLATE, current, STORE, prior=prior, template_dir=FIXTURES)

    assert summary.slides_built == 2
    assert summary.slides_carried == 2
    assert summary.slides_dropped_with_reason == 1
    assert summary.slides_missing_source == 0
    assert summary.as_dict() == {
        "slides_built": 2,
        "slides_carried": 2,
        "slides_dropped_with_reason": 1,
        "slides_missing_source": 0,
    }
    titles = _slide_titles(pptx_bytes)
    assert titles == ["Due Diligence Deck 2025", "Governance"]
    assert len(Presentation(BytesIO(pptx_bytes)).slides) == 2


def test_empty_store_query_fails_unless_allow_empty() -> None:
    manifest = DeckManifest.from_mapping(
        {
            "slides": [
                {
                    "slide_id": "appendix",
                    "title": "Appendix",
                    "layout": "Title and Content",
                    "source": "record:does-not-exist",
                }
            ]
        }
    )

    with pytest.raises(DeckBuildError, match="slide_id 'appendix'"):
        build_deck(TEMPLATE, manifest, STORE, template_dir=FIXTURES)

    allowed = _load_manifest("manifest_allow_empty.json")
    _, summary = build_deck(TEMPLATE, allowed, STORE, template_dir=FIXTURES)
    assert summary.slides_missing_source == 1
    assert summary.slides_built == 1


def test_produced_file_is_valid_presentation_without_office() -> None:
    manifest = _load_manifest("manifest_v1.json")
    pptx_bytes, summary = build_deck(TEMPLATE, manifest, STORE, template_dir=FIXTURES)

    presentation = Presentation(BytesIO(pptx_bytes))
    assert len(presentation.slides) == summary.slides_built == 3
    assert _slide_titles(pptx_bytes) == [
        "Due Diligence Deck",
        "Governance",
        "Risk Overview",
    ]
    assert _slide_bodies(pptx_bytes) == [
        "Synthetic due diligence deck fixture.",
        "Policy approved.",
        "No exceptions & no missing rows.",
    ]


@pytest.mark.parametrize("slide_count", [1, 2])
@pytest.mark.parametrize("as_bytes", [False, True])
def test_populated_template_is_rejected(tmp_path: Path, slide_count: int, as_bytes: bool) -> None:
    presentation = Presentation(TEMPLATE)
    for _ in range(slide_count):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "STALE TEMPLATE SLIDE"
    template = tmp_path / "populated.pptx"
    presentation.save(template)
    original = template.read_bytes()

    with pytest.raises(DeckBuildError, match=rf"template contains {slide_count} existing slide"):
        build_deck(
            original if as_bytes else template,
            _load_manifest("manifest_v1.json"),
            STORE,
            template_dir=FIXTURES,
        )

    assert template.read_bytes() == original


def test_populated_template_is_rejected_before_store_queries(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import deliverable_render.pptx.manifest as manifest_module

    presentation = Presentation(TEMPLATE)
    presentation.slides.add_slide(presentation.slide_layouts[1])
    buffer = BytesIO()
    presentation.save(buffer)

    def unexpected_query(*args: object) -> None:
        pytest.fail("populated template reached store queries before rejection")

    monkeypatch.setattr(manifest_module, "_query_store", unexpected_query)
    with pytest.raises(DeckBuildError, match="template contains 1 existing slide"):
        build_deck(buffer.getvalue(), _load_manifest("manifest_v1.json"), STORE)


@pytest.mark.parametrize(
    "manifest_name", ["manifest_v1.json", "manifest_v2_complete.json", "manifest_allow_empty.json"]
)
def test_successful_deck_contains_exact_active_manifest_slides(manifest_name: str) -> None:
    manifest = _load_manifest(manifest_name)
    content, summary = build_deck(TEMPLATE, manifest, STORE, template_dir=FIXTURES)
    presentation = Presentation(BytesIO(content))

    assert _slide_titles(content) == [
        slide.title for slide in manifest.slides if slide.dropped_reason is None
    ]
    assert len(presentation.slides) == summary.slides_built


def test_empty_manifest_builds_no_slides() -> None:
    content, summary = build_deck(TEMPLATE, DeckManifest(()), STORE)
    assert len(Presentation(BytesIO(content)).slides) == summary.slides_built == 0


def test_unmanifested_slide_added_during_build_is_rejected(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import deliverable_render.pptx.manifest as manifest_module

    original = manifest_module._set_slide_title

    def add_extra_slide(slide: object, title: str) -> None:
        original(slide, title)
        presentation.slides.add_slide(presentation.slide_layouts[1])

    presentation = Presentation(TEMPLATE)
    monkeypatch.setattr(manifest_module, "_load_template", lambda _: presentation)
    monkeypatch.setattr(manifest_module, "_set_slide_title", add_extra_slide)
    with pytest.raises(DeckBuildError, match="slide count.*manifest"):
        build_deck(TEMPLATE, _load_manifest("manifest_v1.json"), STORE, template_dir=FIXTURES)


def test_malformed_manifest_rejects_non_boolean_allow_empty() -> None:
    with pytest.raises(DeckBuildError, match="allow_empty must be a boolean"):
        DeckManifest.from_mapping(
            {
                "slides": [
                    {
                        "slide_id": "appendix",
                        "title": "Appendix",
                        "layout": "Title and Content",
                        "source": "record:gov-24",
                        "allow_empty": "false",
                    }
                ]
            }
        )


def test_malformed_manifest_rejects_non_string_source() -> None:
    with pytest.raises(DeckBuildError, match="source must be a string"):
        DeckManifest.from_mapping(
            {
                "slides": [
                    {
                        "slide_id": "appendix",
                        "title": "Appendix",
                        "layout": "Title and Content",
                        "source": None,
                    }
                ]
            }
        )


def test_static_asset_path_traversal_rejected() -> None:
    manifest = DeckManifest.from_mapping(
        {
            "slides": [
                {
                    "slide_id": "leak",
                    "title": "Leak",
                    "layout": "Title and Content",
                    "source": "static:../cover_blurb.txt",
                }
            ]
        }
    )
    with pytest.raises(DeckBuildError, match="relative path"):
        build_deck(TEMPLATE, manifest, STORE, template_dir=FIXTURES)


def test_deliberate_break_gate_must_fail_then_restore() -> None:
    """Deliberate-break gate: disabling the completeness check must let a bad build through."""
    prior = _load_manifest("manifest_v1.json")
    current = _load_manifest("manifest_v2_incomplete.json")

    original = enforce_completeness

    def noop(_prior: DeckManifest, _current: DeckManifest) -> None:
        return None

    import deliverable_render.pptx.manifest as manifest_module

    manifest_module.enforce_completeness = noop  # type: ignore[assignment]
    try:
        _, summary = build_deck(TEMPLATE, current, STORE, prior=prior, template_dir=FIXTURES)
        assert summary.slides_dropped_with_reason == 0
    finally:
        manifest_module.enforce_completeness = original

    with pytest.raises(DeckBuildError, match="slide_id 'risk'"):
        build_deck(TEMPLATE, current, STORE, prior=prior, template_dir=FIXTURES)


@pytest.mark.parametrize(
    ("slides", "expected"),
    [
        ([], ("gov-24", "gov-25", "risk-25")),
        ([{"source": "record:gov-24"}], ("gov-25", "risk-25")),
        (
            [{"source": "record:gov-24", "dropped_reason": "retired"}],
            ("gov-24", "gov-25", "risk-25"),
        ),
        ([{"source": "static:gov-24"}], ("gov-24", "gov-25", "risk-25")),
        ([{"source": "record:gov-24"}, {"source": "record:gov-24"}], ("gov-25", "risk-25")),
        ([{"source": "record:unknown", "allow_empty": True}], ("gov-24", "gov-25", "risk-25")),
        (
            [
                {"source": "record:gov-25"},
                {"source": "record:risk-25"},
                {"source": "record:gov-24"},
            ],
            (),
        ),
    ],
)
def test_uncovered_record_ids_follow_active_record_sources(
    slides: list[dict[str, object]], expected: tuple[str, ...]
) -> None:
    from deliverable_render.pptx.manifest import uncovered_record_ids

    manifest = DeckManifest.from_mapping(
        {
            "slides": [
                {"slide_id": str(i), "title": "Content", "layout": "Title and Content", **slide}
                for i, slide in enumerate(slides)
            ]
        }
    )
    assert uncovered_record_ids(manifest, STORE) == expected


def test_uncovered_record_ids_empty_store() -> None:
    from deliverable_render.pptx.manifest import uncovered_record_ids

    assert uncovered_record_ids(DeckManifest(()), Store((), ())) == ()
